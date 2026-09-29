#!/usr/bin/env python3
"""Extract full text from PRD and TRD docx files."""
import docx

def extract_docx(path, out_path):
    doc = docx.Document(path)
    lines = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            style = para.style.name if para.style else ""
            if style.startswith("Heading"):
                level = style.replace("Heading ", "")
                try:
                    lvl = int(level)
                    prefix = "#" * lvl
                except ValueError:
                    prefix = "##"
                lines.append(f"{prefix} {text}")
            else:
                lines.append(text)
    # Extract tables
    for i, table in enumerate(doc.tables):
        lines.append(f"\n[TABLE {i+1}]")
        for row in table.rows:
            cells = [c.text.strip().replace("\n", " ") for c in row.cells]
            lines.append(" | ".join(cells))
        lines.append("[/TABLE]\n")
    full = "\n".join(lines)
    with open(out_path, "w", encoding="utf-8") as f:
        f.write(full)
    print(f"=== {path} -> {out_path} ({len(full)} chars) ===")

extract_docx("/home/z/my-project/upload/UniFine_PRD.docx", "/home/z/my-project/scripts/prd_extracted.txt")
extract_docx("/home/z/my-project/upload/UniFine_TRD.docx", "/home/z/my-project/scripts/trd_extracted.txt")
